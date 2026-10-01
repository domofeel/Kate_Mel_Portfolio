from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx"
OUTPUT.mkdir(parents=True, exist_ok=True)

BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(72, 72, 72)
LINK = "0F6F68"
FONT = "Aptos"


def set_font(run, size, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    properties.append(size)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(2.4)
    normal.paragraph_format.line_spacing = 1.04

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(2)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.color.rgb = BLACK
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].font.size = Pt(11.5)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(3)

    if "Role" not in [style.name for style in doc.styles]:
        role = doc.styles.add_style("Role", WD_STYLE_TYPE.PARAGRAPH)
    else:
        role = doc.styles["Role"]
    role.font.name = FONT
    role._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    role._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    role.font.size = Pt(10.1)
    role.font.bold = True
    role.font.color.rgb = BLACK
    role.paragraph_format.space_before = Pt(5)
    role.paragraph_format.space_after = Pt(1)
    role.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = FONT
    bullet._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    bullet.font.size = Pt(9.6)
    bullet.font.color.rgb = BLACK
    bullet.paragraph_format.left_indent = Inches(0.2)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_after = Pt(1.6)
    bullet.paragraph_format.line_spacing = 1.02


def add_header(doc, title, portfolio_url):
    p = doc.add_paragraph(style="Title")
    p.add_run("Ekaterina Melnikova")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(title), 12.2, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run("Novi Sad, Serbia | Residence permit with the right to work | "), 9.5, color=MUTED)
    add_hyperlink(p, "domofeel@gmail.com", "mailto:domofeel@gmail.com")
    set_font(p.add_run(" | "), 9.5, color=MUTED)
    add_hyperlink(p, "LinkedIn", "https://www.linkedin.com/in/kate-mel-71645a23a/")
    set_font(p.add_run(" | "), 9.5, color=MUTED)
    add_hyperlink(p, "Portfolio", portfolio_url)


def add_section(doc, heading, body=None):
    doc.add_paragraph(heading, style="Heading 1")
    if body:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(body), 9.8)


def add_role(doc, role, company, dates, bullets):
    p = doc.add_paragraph(style="Role")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run(f"{role} | {company}"), 10.1, bold=True)
    set_font(p.add_run(f"\t{dates}"), 9.5, bold=True)
    for item in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.keep_together = True
        set_font(p.add_run(item), 9.6)


def add_skills(doc, groups):
    for label, value in groups:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(label + ": "), 9.6, bold=True)
        set_font(p.add_run(value), 9.6)


def add_education(doc):
    add_section(doc, "Education")
    p = doc.add_paragraph()
    set_font(p.add_run("Baltic International Academy, Riga | Master's degree, Marketing in Leisure Industry"), 9.7)


def apply_document_font(doc, font_name):
    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = font_name
            properties = style._element.get_or_add_rPr()
            fonts = properties.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                properties.insert(0, fonts)
            fonts.set(qn("w:ascii"), font_name)
            fonts.set(qn("w:hAnsi"), font_name)

    for fonts in doc.element.body.iter(qn("w:rFonts")):
        fonts.set(qn("w:ascii"), font_name)
        fonts.set(qn("w:hAnsi"), font_name)


def add_product_experience(doc):
    add_section(doc, "Experience")
    add_role(doc, "Senior Product Designer", "Zhelezno", "Jul 2024 - Present", [
        "Led end-to-end product design for B2B dashboards and internal tools, shaping complex workflows into scalable product solutions.",
        "Built and maintained a web and mobile design system, UX patterns and developer documentation, increasing development speed by 50%.",
        "Partnered with Product and Engineering, reducing feature rework by 23% and improving roadmap alignment.",
        "Mentored a junior designer and improved feature delivery time by 14%.",
    ])
    add_role(doc, "Middle+ Product Designer", "Zhelezno", "Jan 2022 - Jul 2024", [
        "Redesigned complex B2B admin dashboards, increasing daily user adoption by 40%.",
        "Conducted user interviews, surveys and usability testing, increasing B2B retention by 42% within 13 months.",
        "Shipped an employee mobile app used daily by 65% of staff, accelerating task completion by 27%.",
    ])
    add_role(doc, "UX/UI Designer", "ServiceHub", "Aug 2021 - Dec 2021", [
        "Improved the usability of the Multifinance fintech SaaS platform through testing and iterative design, increasing employee workflow efficiency by 9%.",
        "Designed a fast-loan banking app from 0 to 1, leading to 30 loan completions within two weeks of launch.",
        "Created branding and a landing page for an IT industry festival, generating 37+ qualified leads.",
    ])
    add_role(doc, "UX/UI Designer", "Apple Concierge", "Oct 2020 - Dec 2021", [
        "Designed B2B delivery and food-service apps adopted by Gazprombank, Sberbank and Alfa-Bank, increasing service value by 23%.",
        "Redesigned KPI dashboards and metric-tracking workflows, reducing manager workflow time by 12%.",
        "Led design from user experiments and wireframes to polished UI and interactive prototypes, improving mobile app user flows by 19%.",
    ])
    add_role(doc, "Web Designer", "CTC Media", "May 2020 - Sep 2020", [
        "Redesigned key B2C and B2B marketplace flows, improving user journeys and supporting adoption among 4K+ active users.",
    ])


def add_ux_ui_experience(doc):
    add_section(doc, "Experience")
    add_role(doc, "UX/UI Designer", "Zhelezno", "Jan 2022 - Now", [
        "Designed complex B2B dashboards and internal tools, turning dense workflows into clear, scalable interfaces.",
        "Built and maintained web and mobile design-system components, UX patterns and developer documentation, increasing development speed by 50%.",
        "Worked with Product and Engineering from interaction concepts and prototypes through handoff and QA, reducing feature rework by 23%.",
        "Mentored a junior designer and improved feature delivery time by 14%.",
        "Redesigned B2B admin dashboards and navigation, increasing daily user adoption by 40%.",
        "Used interviews, surveys and usability testing to refine flows and UI, increasing B2B retention by 42% within 13 months.",
        "Designed an employee mobile app used daily by 65% of staff, accelerating task completion by 27%.",
    ])
    add_role(doc, "UX/UI Designer", "ServiceHub", "Aug 2021 - Dec 2021", [
        "Improved a fintech SaaS platform through usability testing and iterative interface design, increasing employee workflow efficiency by 9%.",
        "Designed a fast-loan banking app from user flows to polished UI, leading to 30 loan completions within two weeks of launch.",
        "Created branding and a landing page for an IT industry festival, generating 37+ qualified leads.",
    ])
    add_role(doc, "UX/UI Designer", "Apple Concierge", "Oct 2020 - Dec 2021", [
        "Designed B2B delivery and food-service mobile apps adopted by Gazprombank, Sberbank and Alfa-Bank, increasing service value by 23%.",
        "Redesigned KPI dashboards and metric-tracking flows, reducing manager workflow time by 12%.",
        "Took work from experiments and wireframes to polished UI and interactive prototypes, improving mobile app user flows by 19%.",
    ])
    add_role(doc, "Web Designer", "CTC Media", "May 2020 - Sep 2020", [
        "Redesigned key B2C and B2B marketplace flows, supporting adoption among 4K+ active users.",
    ])


def build_product_resume():
    doc = Document()
    configure_document(doc)
    add_header(doc, "Product Designer", "https://kate-mel-portfolio-product.vercel.app/")
    add_section(
        doc,
        "Profile",
        "Product Designer focused on B2B SaaS, complex admin tools and mobile products. I connect user research, usability testing and product data with clear workflows and scalable interface solutions, taking features from discovery through design, delivery and iteration.",
    )
    add_section(doc, "Core Skills")
    add_skills(doc, [
        ("Product", "discovery-to-delivery, product strategy, prioritization, metrics and KPIs, roadmap alignment"),
        ("Research and UX", "user interviews, surveys, usability testing, insight synthesis, information architecture, complex workflows"),
        ("Design and delivery", "user flows, wireframes, prototyping, mobile UI, design systems, Figma, developer handoff, documentation"),
    ])
    add_product_experience(doc)
    add_education(doc)
    path = OUTPUT / "Ekaterina_Melnikova_Product_Designer.docx"
    doc.save(path)
    return path


def build_ux_ui_resume():
    doc = Document()
    configure_document(doc)
    add_header(doc, "UX/UI Designer", "https://kate-mel-portfolio-ux-ui.vercel.app/")
    add_ux_ui_experience(doc)
    add_education(doc)
    add_section(
        doc,
        "Skills",
        "UX Design, UI Design, Interaction Design, Information Architecture, User Flows, Wireframing, Interactive Prototyping, User Research, User Interviews, Usability Testing, Responsive Design, Accessibility, Design Systems, Component Libraries, Design Tokens, Developer Handoff, Figma, FigJam, ProtoPie, Adobe Photoshop, Adobe Illustrator, ChatGPT, Claude, Codex.",
    )
    apply_document_font(doc, "Arial")
    path = OUTPUT / "Ekaterina_Melnikova_UXUI_Designer.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_product_resume())
    print(build_ux_ui_resume())
