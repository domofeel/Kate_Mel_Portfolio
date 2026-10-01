from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


root = Path(__file__).resolve().parents[1]
pattern = re.compile(r"(?<!\d)(?:7\+|5\+)|years of experience|—", re.I)

for path in sorted((root / "output" / "pdf").glob("*.pdf")):
    reader = PdfReader(path)
    text = " ".join(page.extract_text() or "" for page in reader.pages)
    links = sum(len(page.get("/Annots", [])) for page in reader.pages)
    print(path.name, "pages", len(reader.pages), "chars", len(text), "links", links, "forbidden", bool(pattern.search(text)))

for path in sorted((root / "output" / "docx").glob("*.docx")):
    document = Document(path)
    text = " ".join(paragraph.text for paragraph in document.paragraphs)
    print(path.name, "paragraphs", len(document.paragraphs), "forbidden", bool(pattern.search(text)))
